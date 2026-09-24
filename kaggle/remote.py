"""Remote-only profile, preflight, validation and downloadable outputs."""
import argparse
from datetime import datetime, timezone
import json
import os
from pathlib import Path
import subprocess
import sys
import zipfile

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'artifacts'
CONFIG = ROOT / 'config/kaggle.yaml'


def configure():
    import yaml
    cfg = yaml.safe_load((ROOT / 'config/pipeline.yaml').read_text(encoding='utf-8'))
    cfg['data_dir'] = str(ROOT / 'data')
    cfg['output_dir'] = str(OUT)
    cfg['extraction']['ocr_use_gpu'] = True
    cfg['extraction']['document_timeout_seconds'] = max(600, cfg['extraction'].get('document_timeout_seconds', 600))
    cfg['retrieval']['embedding_device'] = 'cuda:0'
    cfg['retrieval']['embedding_model_path'] = str(ROOT / '.cache/models/bge-m3')
    cfg['retrieval']['embedding_use_fp16'] = True
    cfg['knowledge']['checklist_mode'] = 'heuristic'
    CONFIG.write_text(yaml.safe_dump(cfg, allow_unicode=True, sort_keys=False), encoding='utf-8')


def preflight():
    import torch
    from docx import Document
    from vn_labor_offline.extract import extract_pdf_docling
    from FlagEmbedding import BGEM3FlagModel
    import faiss, bm25s
    assert torch.cuda.is_available(), 'CUDA unavailable. Enable GPU in Notebook settings.'
    torch.ones(8, device='cuda').sum().item()
    print('Remote GPU:', torch.cuda.get_device_name(0), '; PyTorch:', torch.__version__, flush=True)
    manifest = json.loads((ROOT / 'bundle_manifest.json').read_text(encoding='utf-8'))
    for row in manifest['converted_docs']:
        doc = Document(OUT / '01_extracted/converted_docx' / row['output_name'])
        if not doc.paragraphs and not doc.tables:
            raise RuntimeError(f'Empty converted DOC: {row["output_name"]}')
    # Download/warm models before the per-document timeout begins.
    subprocess.run([sys.executable, str(ROOT / 'scripts/prepare_dense_model.py')], check=True)
    pdf = next((ROOT / 'data').rglob('*.pdf'))
    text = extract_pdf_docling(pdf, ['vi', 'en'], True, 1)
    if not text.strip():
        raise RuntimeError('OCR preflight returned empty text. Check the setup log before running the corpus.')
    print(f'OCR preflight OK: {len(text)} characters on a sample page.', flush=True)
    print('This confirms execution, not legal extraction accuracy.', flush=True)


def refresh_summary():
    from vn_labor_offline.config import load_yaml, resolve_paths
    from vn_labor_offline.util import read_jsonl
    from vn_labor_offline.validation import validate
    cfg = resolve_paths(load_yaml(CONFIG), CONFIG)
    names = ['02_registry/documents.jsonl', '01_extracted/documents.jsonl',
             '03_structure/provisions.jsonl', '03_structure/cases.jsonl',
             '04_knowledge/diagnostic_checklists.jsonl', '05_graph/nodes.jsonl', '05_graph/edges.jsonl']
    if not all((OUT / name).exists() for name in names):
        return
    validate(*(list(read_jsonl(OUT / name)) for name in names), OUT, cfg)


def audit(neo4j=False):
    refresh_summary()
    cmd = [sys.executable, str(ROOT / 'scripts/validate_outputs.py'), '--config', str(CONFIG)]
    if neo4j:
        cmd.append('--neo4j')
    code = subprocess.run(cmd, cwd=ROOT).returncode
    if code not in (0, 1):
        raise RuntimeError(f'Audit process failed: {code}')
    return code


def export():
    reports = OUT / 'reports'
    reports.mkdir(parents=True, exist_ok=True)
    def read(name):
        p = reports / name
        return json.loads(p.read_text(encoding='utf-8')) if p.exists() else {}
    summary = read('summary.json')
    validation = read('final_outputs_validation.json')
    run = read('kaggle_run.json')
    ready = bool(validation.get('ready_for_offline_v1')) and run.get('pipeline_exit_code', 0) == 0
    lines = ['# Báo cáo chạy Kaggle', '', f'Tạo lúc (UTC): {datetime.now(timezone.utc).isoformat()}', '',
             f'OFFLINE v1: **{"PASS" if ready else "CHƯA ĐẠT"}**',
             f'ONLINE-ready: **{"PASS" if validation.get("offline_ready_for_online") else "CHƯA ĐẠT"}**',
             f'Mã thoát pipeline: {run.get("pipeline_exit_code", "chưa chạy trong phiên này")}', '',
             'Mã thoát 0 chỉ xác nhận tiến trình kết thúc; xem validation để đánh giá dữ liệu.', '',
             '| Đầu ra | Trạng thái |', '|---|---|']
    for stage in ('registry', 'structure', 'graph', 'indexes'):
        lines.append(f'| {stage} | {"PASS" if validation.get("stages", {}).get(stage) else "FAIL / chưa kiểm tra"} |')
    lines += ['', '## Số lượng', '']
    for key in ('documents', 'provisions', 'cases', 'diagnostic_items', 'graph_nodes', 'graph_edges'):
        lines.append(f'- {key}: {summary.get(key, "chưa có")}')
    lines += ['', '## Vấn đề còn lại', '']
    lines += [f'- {key}: {value}' for key, value in summary.get('issues_by_type', {}).items()]
    lines += ['', 'Nếu chưa cấu hình Aura, Graph DB chưa được nạp/kiểm chứng nên chưa thể PASS cả bốn đầu ra.',
              'Chi tiết: summary.json, validation_issues.jsonl, final_outputs_validation.md và các log trong reports/.',
              'Báo cáo này được thay mới mỗi lần export; không tự sửa metadata để xóa lỗi.']
    (reports / 'tong_hop_sau_chay.md').write_text('\n'.join(lines), encoding='utf-8')
    archive = ROOT.parent / 'vn_labor_results.zip'
    temporary = archive.with_suffix('.tmp')
    with zipfile.ZipFile(temporary, 'w', compression=zipfile.ZIP_DEFLATED, compresslevel=1) as z:
        for p in sorted(OUT.rglob('*')):
            if p.is_file() and not p.is_symlink() and not any(part.startswith('tmp') for part in p.relative_to(OUT).parts):
                z.write(p, p.relative_to(ROOT).as_posix())
    os.replace(temporary, archive)
    print(f'Download: {archive}; {archive.stat().st_size / 1024**2:.1f} MiB', flush=True)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('action', choices=['configure', 'preflight', 'audit', 'export', 'aura'])
    args = parser.parse_args()
    if args.action == 'configure': configure()
    elif args.action == 'preflight': preflight()
    elif args.action == 'export': export()
    elif args.action == 'audit': sys.exit(audit())
    elif args.action == 'aura':
        from vn_labor_offline.neo4j_loader import load_neo4j
        if not os.getenv('NEO4J_URI', '').startswith('neo4j+s://') or not os.getenv('NEO4J_PASSWORD'):
            raise RuntimeError('Configure Aura credentials using Kaggle Secrets.')
        load_neo4j(OUT)  # Dataset replacement only; never --replace-legacy.
        sys.exit(audit(neo4j=True))


if __name__ == '__main__':
    main()
