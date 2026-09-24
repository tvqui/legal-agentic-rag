"""Regression cases from the September Kaggle extraction review."""
import copy
import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from vn_labor_offline.config import load_yaml, resolve_paths
from vn_labor_offline.extract import extract_one, extract_pdf_pages
from vn_labor_offline.legal_structure import parse_legal_document
from vn_labor_offline.metadata import find_doc_number
from vn_labor_offline.ocr_serialization import reading_lines, legal_plain_text
from vn_labor_offline.segmentation import segment_document

ROOT = Path(__file__).resolve().parents[2]


class StructureTests(unittest.TestCase):
    def test_nested_amendment_quotes_do_not_end_at_inner_quote(self):
        text = ('Điều 1. Sửa đổi\n1. Sửa đổi như sau:\n'
                '“Điều 2. Quy định mới\n1. Khái niệm “người lao động”.\n'
                '2. Nội dung thứ hai.”;\n2. Bãi bỏ quy định cũ.\nĐiều 2. Hiệu lực')
        rows = parse_legal_document({'document_id':'d'}, text)
        self.assertEqual([(p['level'],p['number']) for p in rows],
                         [('ARTICLE','1'),('CLAUSE','1'),('CLAUSE','2'),('ARTICLE','2')])
        self.assertIn('2. Nội dung thứ hai.', rows[1]['text'])

    def test_quoted_points_stay_inside_amending_clause(self):
        text = ('Điều 1. Sửa đổi\n1. Sửa đổi điểm a như sau:\n“a) Nội dung mới;\n'
                'b) Nội dung được trích dẫn.”\n2. Nội dung tiếp theo.')
        rows = parse_legal_document({'document_id':'d'}, text)
        self.assertEqual([p['level'] for p in rows], ['ARTICLE','CLAUSE','CLAUSE'])

    def test_keyword_repair_preserves_original_legal_text(self):
        text = 'Đỉều 1. Quy định\n1. Điều kiện áp dụng:\n- đ) Nội dung của điểm.'
        rows = parse_legal_document({'document_id':'d'}, text)
        self.assertEqual([p['level'] for p in rows], ['ARTICLE','CLAUSE','POINT'])
        self.assertEqual(rows[0]['text'], 'Đỉều 1. Quy định')
        self.assertEqual(rows[2]['number'], 'đ')
        self.assertEqual(rows[2]['text'], '- đ) Nội dung của điểm.')

    def test_ocr_article_keyword_and_number_confusions(self):
        text=('Đíều 9. Việc làm\n1. Nội dung thứ nhất.\n'
              'Điều 1l. Tuyển dụng\n1. Nội dung thứ hai.\n'
              'Điền 12. Quản lý\n1. Nội dung thứ ba.')
        rows=parse_legal_document({'document_id':'d'},text)
        self.assertEqual([p['number'] for p in rows if p['level']=='ARTICLE'],['9','11','12'])
        self.assertEqual(rows[0]['text'],'Đíều 9. Việc làm')

    def test_numbered_annex_title_and_form_are_separate(self):
        parts=segment_document({'document_id':'d'},
            'Điều 1. Quy định\nPhụ lục này áp dụng cho người lao động.\n'
            'Phụ lục I: Danh mục\n1. Dữ liệu biểu mẫu\nMẫu số 01/2026: Báo cáo\n1. Họ tên')
        self.assertEqual([p['segment_type'] for p in parts], ['MAIN_BODY','ANNEX','FORM'])

    def test_resolution_number_without_nq_prefix(self):
        self.assertEqual(find_doc_number('', '326-2016-UBTVQH14_text_copy.doc'), '326/2016/UBTVQH14')


class SerializationTests(unittest.TestCase):
    def test_geometry_orders_words_without_inventing_numbers(self):
        cells=[{'text':'điểm gốc','left':70,'top':10,'right':140,'bottom':20},
               {'text':'Điều 2. Quy định','left':10,'top':30,'right':160,'bottom':40},
               {'text':'đ)','left':10,'top':11,'right':25,'bottom':21}]
        self.assertEqual(reading_lines(cells), 'đ) điểm gốc\nĐiều 2. Quy định')

    def test_serializer_preserves_marker_and_picture_children(self):
        from docling_core.types.doc import DoclingDocument, DocItemLabel
        doc=DoclingDocument(name='legal fixture')
        doc.add_list_item(text='Nội dung gốc',enumerated=True,marker='đ)')
        picture=doc.add_picture()
        doc.add_text(label=DocItemLabel.TEXT,text='Điều 8. Văn bản trong ảnh',parent=picture)
        result=legal_plain_text(doc)
        self.assertIn('đ) Nội dung gốc', result)
        self.assertNotIn('1. ', result)
        self.assertIn('Điều 8. Văn bản trong ảnh', result)


class AttachmentTests(unittest.TestCase):
    def setUp(self):
        import pymupdf as fitz
        self.tmp=tempfile.TemporaryDirectory(); self.root=Path(self.tmp.name)
        self.cfg=resolve_paths(load_yaml(ROOT/'config/pipeline.yaml'),ROOT/'config/pipeline.yaml')
        self.cfg['project_root']=self.root
        self.cfg['extraction']['use_docling']='never'
        (self.root/'config').mkdir(); (self.root/'review/attachments').mkdir(parents=True)
        html=self.root/'landing.html'; html.write_text('Navigation only',encoding='utf-8')
        pdf=self.root/'review/attachments/full.pdf'
        with fitz.open() as doc:
            page=doc.new_page(); page.insert_text((30,30),'Full source PDF body with sufficient native text to avoid the OCR route.')
            doc.save(pdf)
        self.row={'file_id':'f','sha256':hashlib.sha256(html.read_bytes()).hexdigest(),
                  'relative_path':'landing.html','path':str(html),'extension':'.html'}
        self.mapping={'html_sha256':self.row['sha256'],'path':'review/attachments/full.pdf',
                      'sha256':hashlib.sha256(pdf.read_bytes()).hexdigest(),'source_url':'https://example.test/source'}
        self.write_mapping()

    def tearDown(self): self.tmp.cleanup()

    def write_mapping(self):
        (self.root/'config/source_attachments.yaml').write_text(
            json.dumps({'attachments':{'landing.html':self.mapping}}),encoding='utf-8')

    def test_full_text_preserves_source_identity_and_page_hash(self):
        row=extract_one(self.row,self.cfg,self.root/'out')
        self.assertIsNone(row['extraction_error'])
        self.assertIn('Full source PDF body',row['text'])
        self.assertNotIn('Navigation',row['text'])
        self.assertEqual(row['file_id'],self.row['file_id'])
        self.assertEqual(row['sha256'],self.row['sha256'])
        self.assertEqual(row['page_provenance'][0]['sha256'],self.mapping['sha256'])

    def test_sha_bound_attachment_can_replace_scanned_pdf(self):
        import pymupdf as fitz
        source=self.root/'scan.pdf'
        with fitz.open() as doc:
            page=doc.new_page(); page.insert_text((30,30),'scan identity placeholder')
            doc.save(source)
        row={**self.row,'path':str(source),'relative_path':'scan.pdf','extension':'.pdf',
             'sha256':hashlib.sha256(source.read_bytes()).hexdigest()}
        mapping={**self.mapping,'source_sha256':row['sha256']}
        mapping.pop('html_sha256')
        (self.root/'config/source_attachments.yaml').write_text(
            json.dumps({'attachments':{'scan.pdf':mapping}}),encoding='utf-8')
        result=extract_one(row,self.cfg,self.root/'out')
        self.assertIsNone(result['extraction_error'])
        self.assertIn('Full source PDF body',result['text'])
        self.assertEqual(result['sha256'],row['sha256'])
        self.assertEqual(result['text_source']['source_sha256'],row['sha256'])

    def test_docx_header_table_stays_before_body(self):
        from docx import Document
        from vn_labor_offline.extract import extract_docx
        document=Document(); table=document.add_table(rows=1,cols=2)
        table.cell(0,0).text='Luật số: 10/2012/QH13'
        table.cell(0,1).text='Quốc hội'
        document.add_paragraph('Điều 1. Phạm vi điều chỉnh')
        path=self.root/'header.docx'; document.save(path)
        text=extract_docx(path)
        self.assertTrue(text.startswith('Luật số: 10/2012/QH13'))
        self.assertEqual(find_doc_number(text,'10_laodong.doc'),'10/2012/QH13')

    def test_attachment_hash_mismatch_fails_closed(self):
        self.mapping['sha256']='incorrect'; self.write_mapping()
        row=extract_one(self.row,self.cfg,self.root/'out')
        self.assertIn('SHA mismatch',row['extraction_error'])
        self.assertEqual(row['text'],'')

    def test_attachment_cannot_escape_dedicated_directory(self):
        self.mapping['path']='../outside.pdf'; self.write_mapping()
        row=extract_one(self.row,self.cfg,self.root/'out')
        self.assertIn('outside review/attachments',row['extraction_error'])

    def test_reviewed_blank_requires_matching_sha_and_skips_ocr(self):
        review={'sha256':self.row['sha256'],'pages':{1:{'action':'NO_BODY_TEXT','reason':'test fixture review'}}}
        import yaml
        catalog=self.root/'config/page_reviews.yaml'
        catalog.write_text(yaml.safe_dump({'documents':{'landing.html':review}}),encoding='utf-8')
        cfg=copy.deepcopy(self.cfg); cfg['extraction']['use_docling']='always'
        with patch('vn_labor_offline.extract.extract_pdf_docling') as ocr:
            texts,pages=extract_pdf_pages(self.root/'review/attachments/full.pdf',self.row,cfg,self.root/'out')
        ocr.assert_not_called()
        self.assertEqual(texts,['']); self.assertEqual(pages[0]['ocr_status'],'REVIEWED_NO_BODY_TEXT')
        review['sha256']='stale'
        catalog.write_text(yaml.safe_dump({'documents':{'landing.html':review}}),encoding='utf-8')
        with self.assertRaisesRegex(ValueError,'SHA mismatch'):
            extract_pdf_pages(self.root/'review/attachments/full.pdf',self.row,cfg,self.root/'out')

    def test_old_page_cache_does_not_skip_fixed_ocr(self):
        from vn_labor_offline.util import stable_id
        cfg=copy.deepcopy(self.cfg); cfg['extraction']['use_docling']='always'
        cache=self.root/'out/01_extracted/page_cache'; cache.mkdir(parents=True)
        key=stable_id('page-v3',self.row['sha256'],'1',json.dumps(cfg['extraction'],sort_keys=True),prefix='page')
        (cache/(key+'.json')).write_text(json.dumps({'text':'corrupted old OCR'}),encoding='utf-8')
        with patch('vn_labor_offline.extract.extract_pdf_docling',return_value='Correct source text') as ocr:
            pages,_=extract_pdf_pages(self.root/'review/attachments/full.pdf',self.row,cfg,self.root/'out')
        ocr.assert_called_once()
        self.assertEqual(pages,['Correct source text'])

    def test_document_timeout_retries_resume_worker(self):
        from vn_labor_offline.extract import extract_all
        cfg=copy.deepcopy(self.cfg); cfg['extraction']['document_timeout_seconds']=.1
        cfg['extraction']['document_timeout_attempts']=2
        response={'text':'Recovered after page-cache resume','text_chars':33,'extraction_error':None,
                  'page_provenance':[],'extraction_schema':'page-v4-legal-markers'}
        def worker(*args,**kwargs):
            if worker.calls==0:
                worker.calls+=1
                raise __import__('subprocess').TimeoutExpired('worker',.1)
            request=json.loads(Path(args[0][-1]).read_text(encoding='utf-8'))
            Path(request['response']).write_text(json.dumps({**self.row,**response}),encoding='utf-8')
            return __import__('subprocess').CompletedProcess(args[0],0)
        worker.calls=0
        with patch('vn_labor_offline.extract.subprocess.run',side_effect=worker) as run:
            result=extract_all([self.row],cfg,self.root/'out')
        self.assertEqual(run.call_count,2)
        self.assertEqual(result[0]['text'],response['text'])

    def test_checkpoint_reuses_unaffected_document_after_catalog_change(self):
        from vn_labor_offline.extract import extract_all
        cfg=copy.deepcopy(self.cfg)
        response={**self.row,'text':'Existing successful extraction with enough content for reuse.',
                  'text_chars':61,'extraction_error':None,'page_provenance':[],
                  'extraction_schema':'page-v4-legal-markers','text_source':None}

        def worker(*args,**kwargs):
            request=json.loads(Path(args[0][-1]).read_text(encoding='utf-8'))
            Path(request['response']).write_text(json.dumps(response),encoding='utf-8')
            return __import__('subprocess').CompletedProcess(args[0],0)

        # First run creates a successful document cache.
        (self.root/'config/source_attachments.yaml').write_text('{}',encoding='utf-8')
        with patch('vn_labor_offline.extract.subprocess.run',side_effect=worker) as run:
            extract_all([self.row],cfg,self.root/'out')
        self.assertEqual(run.call_count,1)

        # An attachment for another document changes the global catalog hash,
        # but this row's source and review inputs are unchanged.
        (self.root/'config/source_attachments.yaml').write_text(
            json.dumps({'attachments':{'another.pdf':{'source_sha256':'x','path':'review/attachments/x.pdf','sha256':'y'}}}),
            encoding='utf-8')
        with patch('vn_labor_offline.extract.subprocess.run') as run:
            reused=extract_all([self.row],cfg,self.root/'out')
        run.assert_not_called()
        self.assertEqual(reused[0]['text'],response['text'])


if __name__=='__main__': unittest.main()

