"""Rebuild the curated development Gold against the supplied OFFLINE corpus.

No model calls, no writes to the retrieval corpus, no legal approval. The original
v0.1 is retained. All manual edits below are annotations, never pipeline rules.
"""
from __future__ import annotations

import argparse
import copy
import csv
import hashlib
import io
import json
import sys
import zipfile
from collections import Counter
from pathlib import Path

DOC = 'doc_d98ede4425ed3133'
DATE = '2026-09-15'
OLD = 'PhapDienAI_Gold_100_Labor_v0_1.zip'
NEW = 'PhapDienAI_Gold_100_Labor_v0_2.zip'
EXTENDED = 'PhapDienAI_Gold_120_Labor_v0_2.zip'


def digest(data):
    return hashlib.sha256(data).hexdigest()


def dumps(value):
    return json.dumps(value, ensure_ascii=False, indent=2)


def jsonl(rows):
    return ''.join(json.dumps(r, ensure_ascii=False) + '\n' for r in rows)


def build(project: Path, artifact: Path, output: Path):
    sys.path.insert(0, str(project / 'benchmarks/evaluation_suite'))
    from gold import import_candidate, read_candidate
    from types import SimpleNamespace
    original_path = project / 'gold_dataset' / OLD
    originals, old_catalog, old_manifest = read_candidate(original_path)
    with zipfile.ZipFile(artifact) as archive:
        raw_units = archive.read('artifacts/06_indexes/retrieval_units.jsonl')
        units = [json.loads(s) for s in raw_units.decode('utf-8-sig').splitlines() if s.strip()]
        neo = json.loads(archive.read('artifacts/reports/neo4j_validation.json'))
        final = json.loads(archive.read('artifacts/reports/final_outputs_validation.json'))
        dense = json.loads(archive.read('artifacts/reports/dense_validation.json'))
        raw_nodes = archive.read('artifacts/05_graph/nodes.jsonl')
        raw_edges = archive.read('artifacts/05_graph/edges.jsonl')
        graph = [[json.loads(s) for s in raw.decode('utf-8-sig').splitlines() if s.strip()]
                 for raw in (raw_nodes, raw_edges)]
    graph_fp = digest(json.dumps(graph, ensure_ascii=False, sort_keys=True).encode())
    unit_fp = digest(json.dumps(units, ensure_ascii=False, sort_keys=True).encode())
    if graph_fp != neo.get('build_id') or not neo.get('passed'):
        raise ValueError('Corpus graph does not match the validated build')
    if not all(final.get('stages', {}).get(k) is True for k in ('registry','structure','graph','indexes')):
        raise ValueError('Corpus technical stages are incomplete')
    if not dense.get('dense_passed') or dense.get('units') != len(units):
        raise ValueError('Corpus dense validation is incomplete')
    law = [u for u in units if u.get('document_id') == DOC and u.get('kind') == 'PROVISION']
    index = {}
    for u in law:
        key = (str(u.get('article_number') or ''), str(u.get('clause_number') or ''), str(u.get('point_number') or ''))
        if key in index:
            raise ValueError('Ambiguous provision locator: ' + repr(key))
        index[key] = u

    def unit(a, c='', p=''):
        return index[(str(a), str(c), str(p))]

    def key(a, c='', p=''):
        return (str(a), str(c), str(p))

    rows = copy.deepcopy(originals)
    by_id = {r['question_id']: r for r in rows}
    updates = {}

    def edit(qid, reason, question=None, answer=None, refs=None, missing=None):
        r = by_id[qid]
        if question is not None: r['question'] = question
        if answer is not None: r['gold_answer'] = answer
        if refs is not None: updates[qid] = refs
        if missing is not None: r['missing_facts'] = missing
        r['annotation_note'] = reason

    edit('HD001', 'Nêu rõ dữ kiện loại trừ các trường hợp buộc lập văn bản, tránh đồng nhất công việc thông thường với mọi ngoại lệ.',
         question='Tôi 30 tuổi, tự ký hợp đồng làm việc 20 ngày cho công việc thông thường, không phải giúp việc gia đình và không ký qua người được nhóm lao động ủy quyền. Hai bên có thể giao kết bằng lời nói không?')
    edit('HD007', 'Làm rõ lao động thông thường, tránh dùng quy tắc chuyển loại hợp đồng như kết luận vô điều kiện cho nhóm ngoại lệ.',
         question='Tôi 30 tuổi, là công dân Việt Nam, không thuộc nhóm lao động cao tuổi hoặc trường hợp đặc biệt được ký nhiều lần hợp đồng xác định thời hạn. Hợp đồng 24 tháng đã hết hạn, tôi tiếp tục làm việc 40 ngày, hai bên không ký mới. Hợp đồng trở thành loại nào?')
    edit('HD009', 'Câu hỏi yêu cầu quy tắc chung và các ngoại lệ đã được nêu trong chính căn cứ, không suy đoán loại lao động cụ thể.',
         question='Theo quy tắc chung, sau hai hợp đồng xác định thời hạn liên tiếp, nếu người lao động tiếp tục làm việc thì phải ký loại hợp đồng nào? Nêu các nhóm ngoại lệ được dẫn tại điểm c khoản 2 Điều 20.',
         answer='Theo điểm c khoản 2 Điều 20, sau khi ký thêm một hợp đồng xác định thời hạn, nếu tiếp tục làm việc thì phải ký hợp đồng không xác định thời hạn. Ngoại lệ được dẫn gồm người được thuê làm giám đốc trong doanh nghiệp có vốn nhà nước; người lao động cao tuổi theo khoản 1 Điều 149; lao động nước ngoài theo khoản 2 Điều 151; và việc gia hạn đến hết nhiệm kỳ cho thành viên ban lãnh đạo tổ chức đại diện người lao động tại cơ sở theo khoản 4 Điều 177. Ngoại lệ cuối là gia hạn bảo vệ nhiệm kỳ, không phải quyền ký lại vô hạn cho mọi người lao động.',
         refs=[key(20,2,'c'),key(149,1),key(151,2),key(177,4)])
    edit('TV004', 'Loại trừ rõ chức danh quản lý doanh nghiệp, phân biệt yêu cầu công việc với bằng cấp cá nhân.',
         question='Vị trí chuyên môn yêu cầu trình độ cao đẳng, không phải công việc của người quản lý doanh nghiệp. Công ty yêu cầu thử việc 90 ngày có vượt giới hạn không?')
    edit('TV005', 'Giới hạn kết luận vào thời lượng, không suy ra toàn bộ thỏa thuận thử việc hợp pháp.',
         answer='Chỉ xét thời lượng, 45 ngày không vượt giới hạn 60 ngày của khoản 2 Điều 25 đối với công việc này. Điều đó không tự xác nhận mọi nội dung khác của thỏa thuận thử việc đều hợp pháp.')
    edit('TL005', 'Bổ sung công thức đền bù đầy đủ từ chính nguồn; không bịa lãi suất hoặc số tiền khi thiếu dữ kiện.',
         answer='Có. Vì chậm từ 15 ngày trở lên, khoản 4 Điều 97 yêu cầu đền bù ít nhất bằng tiền lãi của số tiền trả chậm, tính theo lãi suất huy động tiền gửi kỳ hạn 01 tháng do ngân hàng nơi công ty mở tài khoản trả lương công bố tại thời điểm trả lương. Chưa có số tiền và lãi suất thực tế để tính một khoản tiền cụ thể.')
    edit('TG009', 'Bổ sung đồng thuận, giới hạn ngày/tháng và thông báo bằng văn bản; không coi ngành dệt may là đủ mọi điều kiện.',
         answer='Có thể, nếu thuộc hoạt động sản xuất, gia công xuất khẩu hàng dệt, may tại điểm a khoản 3 Điều 107; giới hạn năm không quá 300 giờ. Vẫn phải đáp ứng khoản 2 về sự đồng ý và giới hạn ngày, tháng; khi tổ chức theo khoản 3 phải thông báo bằng văn bản theo khoản 4 Điều 107. Không phải cứ doanh nghiệp dệt may là tự động được bỏ các điều kiện này.',
         refs=[key(107,3),key(107,3,'a'),key(107,2,'a'),key(107,2,'b'),key(107,4)])
    edit('TG014', 'Kết hợp phần mở đầu Điều 108 với khoản 2 để chứng minh quyền vượt giới hạn, vẫn giữ ngoại lệ an toàn.',
         answer='Có thể yêu cầu làm thêm trong trường hợp bảo vệ tính mạng con người, phòng ngừa hoặc khắc phục thiên tai theo phần mở đầu và khoản 2 Điều 108, không bị giới hạn số giờ tại Điều 107. Tuy nhiên, vẫn giữ ngoại lệ có nguy cơ ảnh hưởng tính mạng, sức khỏe người lao động theo pháp luật an toàn, vệ sinh lao động.',
         refs=[key(108),key(108,2)])
    edit('NP005', 'Loại trừ nhóm hưởng 14 ngày và thâm niên; hỏi số cụ thể trong đúng một kỳ 12 tháng.',
         question='Tôi 30 tuổi, không phải người khuyết tật, làm công việc trong điều kiện bình thường và vừa đủ 12 tháng cho cùng một công ty. Trong kỳ 12 tháng đó tôi được hưởng bao nhiêu ngày nghỉ hằng năm có lương?',
         answer='Theo điểm a khoản 1 Điều 113, được nghỉ hằng năm 12 ngày làm việc và hưởng nguyên lương theo hợp đồng lao động. Dữ kiện mới đủ 12 tháng chưa phát sinh ngày tăng thêm do đủ 5 năm theo Điều 114; không tự cộng thêm thâm niên không được cung cấp.',
         refs=[key(113,1,'a'),key(114)])
    edit('NP006', 'Hỏi mức cơ sở theo nhóm công việc để không vô tình bỏ ngày tăng thêm do thâm niên.',
         question='Người làm nghề nặng nhọc, độc hại, nguy hiểm đủ 12 tháng cho một người sử dụng lao động được hưởng mức nghỉ hằng năm cơ sở bao nhiêu ngày làm việc, chưa tính ngày tăng thêm do thâm niên?')
    edit('NP007', 'Hỏi mức cơ sở, không giả định thâm niên bằng 0.',
         question='Người làm nghề đặc biệt nặng nhọc, độc hại, nguy hiểm đủ 12 tháng cho một người sử dụng lao động được hưởng mức nghỉ hằng năm cơ sở bao nhiêu ngày làm việc, chưa tính ngày tăng thêm do thâm niên?')
    for qid, months, days in [('CD001',None,'45 ngày'),('CD002',24,'30 ngày'),('CD003',9,'03 ngày làm việc')]:
        contract = 'không xác định thời hạn' if months is None else f'xác định thời hạn {months} tháng'
        edit(qid, 'Cung cấp đủ dữ kiện quyết định thời hạn báo trước; không ép chatbot hỏi lại điều đã biết.',
             question=f'Tôi làm theo hợp đồng {contract}, nghỉ vì lý do cá nhân, không thuộc trường hợp được nghỉ không cần báo trước và không làm ngành, nghề, công việc đặc thù. Tôi phải báo trước tối thiểu bao lâu?',
             answer=f'Phải báo trước ít nhất {days}, theo '+{'CD001':'điểm a','CD002':'điểm b','CD003':'điểm c'}[qid]+' khoản 1 Điều 35. Đây là thời hạn cho dữ kiện đã loại trừ ngành nghề đặc thù và quyền nghỉ không cần báo trước.')
    edit('CD007', 'Bổ sung hợp đồng, tình trạng được bảo vệ và căn cứ báo trước, không chỉ dẫn lý do chấm dứt.',
         question='Nhân viên có hợp đồng 24 tháng, không làm công việc đặc thù và không thuộc trường hợp được bảo vệ tại Điều 37. Người này thường xuyên không hoàn thành công việc theo tiêu chí đánh giá hợp lệ, được ban hành sau khi tham khảo ý kiến tổ chức đại diện người lao động tại cơ sở nơi có tổ chức này. Công ty có quyền đơn phương chấm dứt không và phải báo trước bao lâu?',
         answer='Có thể thực hiện quyền đơn phương theo điểm a khoản 1 Điều 36 nếu tình trạng thường xuyên không hoàn thành công việc thật sự đáp ứng tiêu chí hợp lệ. Với hợp đồng 24 tháng và dữ kiện đã loại trừ nhóm đặc thù, phải báo trước ít nhất 30 ngày theo điểm b khoản 2 Điều 36. Không suy từ một lần không hoàn thành công việc thành điều kiện thường xuyên.',
         refs=[key(36,1,'a'),key(36,2,'b')])
    edit('CD008', 'Tách quyền không báo trước với thủ tục sa thải; đã loại trừ trường hợp bảo vệ.',
         question='Người lao động không thuộc các trường hợp bảo vệ tại Điều 37, tự ý bỏ việc không có lý do chính đáng đủ 5 ngày làm việc liên tục. Nếu công ty đơn phương chấm dứt theo điểm e khoản 1 Điều 36 thì có phải báo trước không?')
    edit('CD009', 'Bổ sung quy định trực tiếp cấm chấm dứt vì mang thai; khoản 3 Điều 37 vẫn là căn cứ bảo vệ.',
         answer='Không được đơn phương chấm dứt chỉ vì người lao động mang thai. Khoản 3 Điều 37 quy định trường hợp người sử dụng lao động không được thực hiện quyền đơn phương; khoản 3 Điều 137 còn trực tiếp cấm sa thải hoặc đơn phương chấm dứt vì lý do mang thai, với các trường hợp ngoại lệ về tình trạng của người sử dụng lao động được luật nêu. Không đồng nhất hết hạn hợp đồng với đơn phương chấm dứt.',
         refs=[key(37,3),key(137,3)])
    edit('CD012', 'Nêu rõ các nhóm nghĩa vụ trong nguồn thay cho cụm quyền lợi liên quan quá chung; giữ điều kiện phát sinh.',
         answer='Có. Theo khoản 1 Điều 41, công ty phải nhận người lao động trở lại, trả tiền lương và đóng bảo hiểm xã hội, bảo hiểm y tế, bảo hiểm thất nghiệp cho những ngày không được làm việc, đồng thời trả thêm ít nhất 02 tháng tiền lương theo hợp đồng. Nếu vi phạm thời hạn báo trước tại khoản 2 Điều 36 thì còn trả khoản tương ứng ngày thiếu báo trước. Người lao động hoàn trả trợ cấp thôi việc, mất việc đã nhận nếu có; nếu không còn vị trí cũ mà vẫn muốn làm việc thì hai bên thỏa thuận sửa đổi hợp đồng.')
    edit('TD001', 'Căn cứ cho các nhánh chính, đặc thù và quyền không báo trước đều có thật; ngày tra cứu đã được cung cấp riêng.',
         refs=[key(35,1,'a'),key(35,1,'b'),key(35,1,'c'),key(35,1,'d'),key(35,2)],
         missing=['Loại và thời hạn hợp đồng','Ngành, nghề, công việc có thuộc nhóm đặc thù','Lý do nghỉ và tình tiết được nghỉ không cần báo trước'])
    edit('TD002', 'Thiếu căn cứ nhóm quản lý trong v0.1; bổ sung khoản 1 Điều 25.',
         refs=[key(25),key(25,1),key(25,2),key(25,3),key(25,4)])
    edit('TD003', 'Số 6 triệu tự nó không buộc phải biết vùng trong mọi cách hiểu. Làm rõ đây là tổng thu nhập có cả tăng ca/thưởng, chưa biết lương công việc và giờ bình thường; không dùng ngưỡng ngoài qrels.',
         question='Công ty thông báo tổng thu nhập của tôi khoảng 6 triệu đồng/tháng, gồm cả tiền tăng ca và thưởng nhưng chưa tách lương công việc. Chỉ dựa vào số này có kết luận đã đáp ứng mức lương tối thiểu không?',
         answer='Chưa thể kết luận chỉ từ tổng thu nhập gồm cả tăng ca và thưởng. Cần biết phần lương theo công việc hoặc chức danh, thời giờ làm việc bình thường và cách trả lương, cùng địa bàn áp dụng để đối chiếu mức tối thiểu tương ứng. Khoản 2 Điều 90 yêu cầu mức lương theo công việc hoặc chức danh không thấp hơn mức tối thiểu; khoản 2 Điều 91 xác lập mức tối thiểu theo vùng, tháng hoặc giờ. Không lấy tổng thu nhập 6 triệu để mặc nhiên chứng minh phần lương công việc đã đạt mức tối thiểu.',
         missing=['Lương theo công việc hoặc chức danh, tách tiền tăng ca và thưởng','Địa bàn áp dụng','Thời giờ làm việc bình thường và hình thức trả lương'])
    edit('TD004', 'Bổ sung điều khoản tỷ lệ cho người làm chưa đủ năm, khuyết tật/tuổi cũng là dữ kiện phân nhóm.',
         refs=[key(113,1,'a'),key(113,1,'b'),key(113,1,'c'),key(113,2),key(114)],
         missing=['Số tháng và thâm niên làm việc cho cùng người sử dụng lao động','Nhóm điều kiện công việc','Độ tuổi và tình trạng khuyết tật nếu liên quan phân nhóm'])
    edit('TD006', 'Bổ sung định nghĩa trái luật, phạm vi báo trước và hậu quả; không dùng phần mở đầu một khoản như toàn bộ quy tắc.',
         refs=[key(35,1,'a'),key(35,1,'b'),key(35,1,'c'),key(35,1,'d'),key(35,2),key(39),key(40,2)],
         missing=['Loại và thời hạn hợp đồng','Lý do nghỉ; trường hợp không cần báo trước','Số ngày đã báo trước','Có thuộc công việc đặc thù','Công ty có thỏa thuận cho nghỉ sớm hay không'])
    edit('TD007', 'Phân biệt xác định quyền hưởng với tính mức, không chỉ dùng thời gian đóng bảo hiểm làm điều kiện duy nhất.',
         answer='Cần biết trường hợp chấm dứt hợp đồng, thời gian làm việc thường xuyên và có thuộc diện hưởng lương hưu hoặc trường hợp loại trừ ở Điều 46 hay không để xác định quyền hưởng. Để tính mức còn cần thời gian đã tham gia bảo hiểm thất nghiệp, thời gian đã được chi trả trợ cấp và tiền lương bình quân theo khoản 3 Điều 46. Không thể kết luận cứ nghỉ việc là được trợ cấp.',
         refs=[key(46,1),key(46,2),key(46,3)],
         missing=['Trường hợp chấm dứt','Thời gian làm việc thường xuyên','Có thuộc trường hợp loại trừ','Thời gian bảo hiểm thất nghiệp và trợ cấp đã nhận; lương tính trợ cấp nếu cần tính mức'])
    edit('TD009', 'Câu cũ có thể trả lời quy tắc chung mà không cần mọi dữ kiện. Đổi sang yêu cầu tính số tiền cụ thể để nhãn thiếu dữ kiện có cơ sở.',
         question='Tôi đi làm vào một ngày nghỉ lễ và có cả giờ làm ban đêm. Công ty phải trả cho tôi chính xác bao nhiêu tiền cho ngày đó?',
         answer='Chưa tính được số tiền cụ thể. Cần biết tiền lương hoặc đơn giá của công việc, số giờ làm ban ngày/ban đêm và hình thức hưởng lương. Khoản 1 điểm c Điều 98 đặt mức làm thêm ngày lễ ít nhất 300%; khoản 2 và khoản 3 quy định các khoản làm đêm và làm thêm ban đêm. Không được chỉ trả lời 300% rồi bỏ các khoản làm đêm, hoặc bịa số tiền khi chưa có giờ và đơn giá.',
         refs=[key(98,1,'c'),key(98,2),key(98,3)],
         missing=['Tiền lương hoặc đơn giá công việc','Số giờ làm ban ngày và ban đêm','Hình thức hưởng lương'])
    edit('TD012', 'Thay căn cứ tiêu đề Điều 14 bằng hai khoản chứa quy tắc thực chất.',
         refs=[key(14,1),key(14,2)])
    future = [copy.deepcopy(by_id[qid]) for qid in ['TD013','TD014','TD015']]
    for qid in ['TD013','TD014','TD015']:
        by_id[qid]['expected_status'] = 'ANSWERABLE'
        by_id[qid]['missing_facts'] = []
        by_id[qid]['category'] = 'Nghỉ phép và nghỉ lễ' if qid=='TD014' else 'Tiền lương'
    edit('TD013', 'Thay câu yêu cầu con số năm 2028 không có trong artifact bằng câu quy tắc được corpus hỗ trợ; bản gốc giữ riêng cho phép thử từ chối.',
         question='Theo Bộ luật Lao động, cơ quan nào quyết định và công bố mức lương tối thiểu, dựa trên khuyến nghị của cơ quan nào?',
         answer='Theo khoản 4 Điều 91, Chính phủ quyết định và công bố mức lương tối thiểu trên cơ sở khuyến nghị của Hội đồng tiền lương quốc gia.',refs=[key(91,4)])
    edit('TD014', 'Thay yêu cầu lịch dương năm 2028 chưa có nguồn bằng quy tắc Quốc khánh và thẩm quyền chọn ngày có trong corpus.',
         question='Theo Bộ luật Lao động, người lao động Việt Nam được nghỉ Quốc khánh mấy ngày hưởng nguyên lương? Ngoài ngày 2 tháng 9 thì ai quyết định ngày nghỉ còn lại?',
         answer='Điểm đ khoản 1 Điều 112 quy định Quốc khánh nghỉ 02 ngày hưởng nguyên lương: ngày 02 tháng 9 dương lịch và 01 ngày liền kề trước hoặc sau. Theo khoản 3 Điều 112, hằng năm Thủ tướng Chính phủ quyết định cụ thể ngày nghỉ này; quy tắc khung không tự xác định lịch cụ thể của một năm chưa có quyết định.',
         refs=[key(112,1,'đ'),key(112,3)])
    edit('TD015', 'Thay dự báo sửa luật năm 2029 bằng quy tắc các khoản làm thêm ban đêm đã có trong corpus, không cộng các tỷ lệ khác cơ sở một cách máy móc.',
         question='Khi người lao động làm thêm giờ vào ban đêm, ngoài tiền làm thêm giờ và khoản làm đêm 30%, luật còn quy định trả thêm khoản nào? Nêu rõ cơ sở tính, không cần tính thành số tiền.',
         answer='Theo khoản 3 Điều 98, ngoài tiền trả theo khoản 1 và khoản 2, còn được trả thêm 20% tiền lương tính theo đơn giá hoặc tiền lương theo công việc làm vào ban ngày của ngày làm việc bình thường, ngày nghỉ hằng tuần hoặc ngày nghỉ lễ, tết tương ứng. Không cộng máy móc 20% vào mọi tổng tỷ lệ nếu cơ sở tính khác nhau.',
         refs=[key(98,3),key(98,2),key(98,1)])

    # Optional additions broaden legal topics and multi-task reasoning. They are
    # separate from the default 100 so an existing Kaggle budget is not raised.
    additions = [
        ('MR001','Kỷ luật','Công ty phạt tôi 500 nghìn đồng vì đi muộn thay cho xử lý kỷ luật lao động. Cách phạt này có được phép không?', 'Không. Khoản 2 Điều 127 cấm phạt tiền, cắt lương thay việc xử lý kỷ luật lao động. Không nhầm hình thức phạt này với bồi thường thiệt hại có căn cứ riêng.',[key(127,2)]),
        ('MR002','Kỷ luật','Cùng một hành vi vi phạm kỷ luật, công ty vừa khiển trách vừa kéo dài thời hạn nâng lương. Có được áp dụng đồng thời không?', 'Không. Khoản 2 Điều 122 không cho áp dụng nhiều hình thức xử lý kỷ luật lao động đối với một hành vi vi phạm.',[key(122,2)]),
        ('MR003','Kỷ luật','Tôi đang nghỉ ốm có sự đồng ý của công ty. Công ty có được xử lý kỷ luật tôi ngay trong thời gian nghỉ này không?', 'Không được xử lý trong thời gian này theo điểm a khoản 4 Điều 122. Đây là giới hạn về thời điểm xử lý, không có nghĩa hành vi vi phạm đương nhiên được xóa bỏ.',[key(122,4,'a')]),
        ('MR004','Kỷ luật','Thời hạn tạm đình chỉ công việc thông thường và trường hợp đặc biệt tối đa là bao lâu? Trong thời gian đó được tạm ứng mấy phần trăm tiền lương?', 'Theo khoản 2 Điều 128, thông thường không quá 15 ngày, trường hợp đặc biệt không quá 90 ngày; được tạm ứng 50% tiền lương trước khi đình chỉ. Hết thời hạn, người sử dụng lao động phải nhận người lao động trở lại làm việc.',[key(128,2)]),
        ('MR005','Kỷ luật','Tôi không bị xử lý kỷ luật sau thời gian tạm đình chỉ công việc. Công ty có phải trả đủ lương của thời gian bị đình chỉ không?', 'Có. Khoản 4 Điều 128 yêu cầu trả đủ tiền lương cho thời gian bị tạm đình chỉ khi người lao động không bị xử lý kỷ luật.',[key(128,4)]),
        ('MR006','Nội quy','Công ty sử dụng 12 người lao động. Nội quy lao động có bắt buộc bằng văn bản không?', 'Có. Khoản 1 Điều 118 yêu cầu người sử dụng lao động phải ban hành nội quy, và từ 10 người lao động trở lên thì nội quy phải bằng văn bản.',[key(118,1)]),
        ('MR007','Lao động chưa thành niên','Một người lao động 14 tuổi có được bố trí làm 5 giờ mỗi ngày hoặc làm thêm ban đêm không? Chỉ đánh giá giới hạn thời giờ, không đánh giá danh mục nghề được tuyển.', 'Không. Khoản 1 Điều 146 quy định người chưa đủ 15 tuổi làm không quá 4 giờ/ngày, 20 giờ/tuần; không được làm thêm giờ hoặc làm việc ban đêm.',[key(146,1)]),
        ('MR008','Lao động chưa thành niên','Người lao động 16 tuổi được bố trí thời giờ làm việc bình thường 8 giờ/ngày, 48 giờ/tuần. Riêng giới hạn theo tuần có đúng không?', 'Không. Khoản 2 Điều 146 giới hạn người từ đủ 15 đến chưa đủ 18 tuổi không quá 8 giờ/ngày và 40 giờ/tuần. 48 giờ/tuần vượt mức này.',[key(146,2)]),
        ('MR009','Lao động nữ','Trong thời gian nuôi con dưới 12 tháng tuổi, lao động nữ được nghỉ mỗi ngày bao nhiêu phút trong giờ làm việc và thời gian ấy có hưởng lương không?', 'Khoản 4 Điều 137 quy định nghỉ 60 phút mỗi ngày trong thời gian làm việc và vẫn hưởng đủ tiền lương theo hợp đồng lao động.',[key(137,4)]),
        ('MR010','Lao động nữ','Lao động nữ trong thời gian hành kinh được nghỉ mỗi ngày tối thiểu bao nhiêu phút theo Bộ luật Lao động? Thời gian này có được trả lương không?', 'Theo khoản 4 Điều 137, được nghỉ mỗi ngày 30 phút trong thời gian làm việc và vẫn hưởng đủ tiền lương theo hợp đồng lao động.',[key(137,4)]),
        ('MR011','Lao động cao tuổi','Người lao động cao tuổi và doanh nghiệp có được thỏa thuận ký nhiều lần hợp đồng xác định thời hạn không?', 'Có. Khoản 1 Điều 149 cho phép hai bên thỏa thuận giao kết nhiều lần hợp đồng xác định thời hạn khi sử dụng người lao động cao tuổi.',[key(149,1)]),
        ('MR012','Lao động nước ngoài','Giấy phép lao động của người lao động nước ngoài còn hiệu lực 18 tháng. Có được ký hợp đồng lao động 24 tháng vượt quá thời hạn giấy phép đó không?', 'Không. Khoản 2 Điều 151 quy định thời hạn hợp đồng lao động không được vượt quá thời hạn giấy phép lao động. Không kết luận hợp đồng 24 tháng hợp lệ chỉ vì chưa vượt 36 tháng.',[key(151,2)]),
        ('MR013','Đào tạo','Công ty dùng kinh phí của mình để đào tạo nâng cao kỹ năng cho người lao động. Hợp đồng đào tạo nghề có phải lập thành hai bản, mỗi bên giữ một bản không?', 'Có. Khoản 1 Điều 62 yêu cầu ký hợp đồng đào tạo nghề trong trường hợp đào tạo nâng cao từ kinh phí người sử dụng lao động, kể cả tài trợ của đối tác; hợp đồng làm thành 02 bản, mỗi bên giữ 01 bản.',[key(62,1)]),
        ('MR014','Tranh chấp','Tranh chấp lao động cá nhân về việc bị sa thải có bắt buộc hòa giải qua hòa giải viên lao động trước khi yêu cầu tòa án giải quyết không? Chỉ hỏi điều kiện hòa giải trước, không hỏi toàn bộ thủ tục khởi kiện.', 'Không bắt buộc đối với loại tranh chấp này. Điểm a khoản 1 Điều 188 liệt kê tranh chấp về xử lý kỷ luật sa thải hoặc đơn phương chấm dứt trong các ngoại lệ không bắt buộc qua thủ tục hòa giải.',[key(188,1),key(188,1,'a')]),
        ('MR015','Tranh chấp','Một tranh chấp lao động cá nhân không có sự kiện bất khả kháng, trở ngại khách quan hoặc lý do khác làm dừng tính thời hiệu. Thời hiệu yêu cầu Tòa án giải quyết là bao lâu và tính từ khi nào?', 'Theo khoản 3 Điều 190, thời hiệu là 01 năm kể từ ngày phát hiện hành vi mà bên tranh chấp cho rằng quyền và lợi ích hợp pháp của mình bị vi phạm. Không tự tính từ ngày ký hợp đồng hoặc mặc định ngày nghỉ việc.',[key(190,3)]),
        ('MR016','Nghỉ phép và nghỉ lễ','Khi nghỉ hằng năm, tôi đi đường bộ cả đi và về hết 4 ngày, đây là lần nghỉ duy nhất có tính thời gian đi đường trong năm. Có được tính thêm thời gian ngoài phép năm không?', 'Theo khoản 6 Điều 113, nếu đi đường bộ, đường sắt, đường thủy cả đi và về trên 2 ngày thì từ ngày thứ 3 trở đi tính thêm ngoài ngày nghỉ hằng năm, chỉ cho 01 lần nghỉ trong năm. Với 4 ngày đi đường, có 2 ngày tính thêm.',[key(113,6)]),
        ('MR017','Nghỉ việc riêng','Cha nuôi của vợ tôi qua đời. Tôi được nghỉ việc riêng hưởng nguyên lương bao nhiêu ngày và có phải thông báo công ty không?', 'Được nghỉ 03 ngày hưởng nguyên lương theo điểm c khoản 1 Điều 115 và phải thông báo người sử dụng lao động theo phần mở đầu khoản 1.',[key(115,1),key(115,1,'c')]),
        ('MR018','Quyền và nghĩa vụ','Trước khi giao kết hợp đồng, công ty có nghĩa vụ cung cấp trung thực thông tin về công việc, tiền lương và bảo hiểm cho tôi không?', 'Có. Khoản 1 Điều 16 yêu cầu người sử dụng lao động cung cấp trung thực các thông tin đó cùng những nội dung khác liên quan trực tiếp đến giao kết hợp đồng được luật quy định.',[key(16,1)]),
        ('MR019','Đa yêu cầu','Tôi 30 tuổi, không phải người khuyết tật, vừa làm đủ 12 tháng cho cùng một công ty trong điều kiện bình thường. Công ty giữ bản chính bằng đại học và quy định kỳ nghỉ hằng năm của 12 tháng đó chỉ có 6 ngày hưởng lương. Hãy kết luận riêng việc giữ bằng và số ngày phép tối thiểu.', 'Vấn đề 1: không được giữ bản chính bằng đại học theo khoản 1 Điều 17. Vấn đề 2: 6 ngày không đạt mức 12 ngày làm việc nghỉ hằng năm hưởng nguyên lương theo điểm a khoản 1 Điều 113 cho dữ kiện này. Phải trả lời cả hai vấn đề, không lấy căn cứ của vấn đề này chứng minh vấn đề kia.',[key(17,1),key(113,1,'a')]),
        ('MR020','Đa yêu cầu','Tôi thử việc lần đầu ở vị trí chuyên môn yêu cầu cao đẳng, không phải người quản lý doanh nghiệp; hợp đồng không thuộc loại dưới một tháng. Công ty yêu cầu thử việc 90 ngày và trả 80% lương của công việc. Hãy đánh giá riêng thời gian và mức lương.', 'Thời gian: 90 ngày vượt mức tối đa 60 ngày theo khoản 2 Điều 25. Tiền lương: 80% thấp hơn mức tối thiểu 85% lương công việc theo Điều 26. Hai kết luận cần căn cứ riêng; không trả lời một phần rồi coi cả câu đã hoàn tất.',[key(25,2),key(26)]),
    ]
    extra = []
    for qid, category, question, answer, refs in additions:
        r = copy.deepcopy(rows[0])
        r.update(question_id=qid, category=category, question=question, gold_answer=answer,
                 expected_status='ANSWERABLE', missing_facts=[],
                 annotation_note='Câu bổ sung phát triển; căn cứ đã đối chiếu trực tiếp artifact, chưa phải nhãn chuyên gia hoặc tập kiểm thử độc lập.')
        extra.append(r)
        updates[qid] = refs

    old_by_id = {r['question_id']: r for r in originals}
    unit_by_id = {u['unit_id']: u for u in units}

    def loc(u):
        return 'Điều ' + u['article_number'] + ((' khoản ' + u['clause_number']) if u['clause_number'] else '') + ((' điểm ' + u['point_number']) if u['point_number'] else '')

    def finalize(r):
        r['applicable_date'] = DATE
        r['temporal_verification_status'] = 'CORPUS_BOUND_NOT_HUMAN_VERIFIED'
        r['verification_status'] = 'DRAFT_NEEDS_HUMAN_LEGAL_REVIEW'
        r['evaluation_scope'] = 'CORPUS_SUPPORTED_DEVELOPMENT_NOT_HELD_OUT'
        r['original_question_id'] = r['question_id'] if r['question_id'] in old_by_id else None
        refs = [unit(*k) for k in updates[r['question_id']]] if r['question_id'] in updates else [unit_by_id[i] for i in r['gold_evidence_ids']]
        r['gold_evidence_ids'] = [u['unit_id'] for u in refs]
        r['gold_evidence_locators'] = [loc(u) for u in refs]
        r['evidence_details'] = [{'evidence_id':u['unit_id'], 'locator':loc(u),
             'evidence_excerpt':u['source_text'], 'ancestor_context':u.get('ancestor_context') or ''} for u in refs]
        r['source_document_id'] = DOC
        r['source_url'] = refs[0]['source_url']
        r['source_support_status'] = 'ALL_ANNOTATED_UNITS_PRESENT_WITH_DATED_CORPUS_CONTEXT'
        # Helpful scoring instructions are not a requirement to repeat a fixed
        # answer verbatim or to accept legally incomplete model output.
        r['scoring_note'] = 'Chấp nhận diễn đạt tương đương, điều kiện/ngoại lệ đúng và căn cứ tương đương đã duyệt; không dùng khớp trạng thái thay cho đánh giá nội dung. Trả lời có cảnh báo provisional không tự động là sai.'
        return r

    rows = [finalize(r) for r in rows]
    extra = [finalize(r) for r in extra]
    change_log = []
    for r in rows:
        old = old_by_id[r['question_id']]
        fields = [k for k in ('category','question','gold_answer','gold_evidence_ids','expected_status','missing_facts','applicable_date') if r[k] != old[k]]
        change_log.append({'question_id':r['question_id'],'changed_fields':fields,
            'reason':r.get('annotation_note') or 'Đọc đối chiếu câu hỏi/đáp án và căn cứ trong corpus; ghim ngày đo về snapshot thực tế.',
            'before':{k:old[k] for k in fields}, 'after':{k:r[k] for k in fields},
            'review_kind':'ENGINEERING_CONTENT_REVIEW_NOT_EXPERT_APPROVAL'})
    output.mkdir(parents=True, exist_ok=True)
    validations = {}
    for name, selected in [(NEW,rows),(EXTENDED,rows+extra)]:
        evidence_ids = sorted({i for r in selected for i in r['gold_evidence_ids']})
        catalog = []
        for uid in evidence_ids:
            u = unit_by_id[uid]
            catalog.append({'evidence_id':uid,'source_document_id':u['document_id'],
                'document_number':u['document_number'],'instrument_number':u['instrument_number'],
                'locator':loc(u),'evidence_text':u['source_text'],
                'ancestor_context':u.get('ancestor_context') or '',
                'source_text_sha256':digest(u['source_text'].encode()),
                'source_url':u['source_url'],'is_in_retrieval_units':True,
                'temporal_verified_at_provision':bool(u.get('provision_temporal_verified')),
                'provenance_status':u.get('provenance_status'), 'provenance_span':u.get('provenance_span'),
                'review_status':'NEEDS_HUMAN_LEGAL_REVIEW'})
        manifest = {**old_manifest,'dataset_name':name.removesuffix('.zip'),
            'created_at':'2026-10-30','data_cutoff':DATE,'size':len(selected),
            'counts_by_category':dict(Counter(r['category'] for r in selected)),
            'counts_by_expected_status':dict(Counter(r['expected_status'] for r in selected)),
            'distinct_evidence_ids':len(catalog),'source_parent_zip':OLD,
            'source_parent_zip_sha256':digest(original_path.read_bytes()),
            'corpus_zip':'vn_labor_results.zip (result_28-10)',
            'corpus_binding':{'build_id':graph_fp,'retrieval_unit_fingerprint':unit_fp},
            'review_status':'DRAFT - corpus-aligned engineering review, no human legal approval',
            'use_restriction':'Development comparison only; edited using prior failures, not independent hold-out or proof of complete legal coverage. No prediction of future law.',
            'coverage_scope':'One consolidated Labor Code; optional 20 additions broaden topics, but not every document, historical period, case law, conversation or every legal issue.',
            'temporal_limitation':'All core query dates deliberately fixed to corpus snapshot 2026-09-15, not represented as present-day legal advice. Provision intervals and sources remain provisional.'}
        files = {'labor_gold_100.jsonl':jsonl(selected), 'evidence_catalog.jsonl':jsonl(catalog),
                 'manifest.json':dumps(manifest), 'change_log.json':dumps(change_log),
                 'future_challenges_NOT_IN_MAIN_SCORE.jsonl':jsonl(future),
                 'README.md': '# Gold ứng viên v0.2\n\nChỉ đo trên snapshot 15/09/2026 của result_28-10. Không giải nén vào corpus.\nCác nhãn/đáp án chỉ dùng để chấm sau khi hỏi, không đưa vào prompt hoặc retrieval.\nBản 100 có 88 câu có căn cứ để trả lời và 12 câu cần hỏi thêm; bản 120 bổ sung 20 câu.\nTên member labor_gold_100.jsonl được giữ vì tương thích reader; size trong manifest là số câu thật.\nBa câu tương lai gốc giữ trong file NOT_IN_MAIN_SCORE, không dùng làm qrels dương tính.\nKhông phải Gold được chuyên gia duyệt, không phải hold-out, không chứng minh model luôn trả lời được.\n',
                 'labor_gold.csv':csv_text(selected)}
        stream = io.BytesIO()
        with zipfile.ZipFile(stream,'w',zipfile.ZIP_DEFLATED) as z:
            for filename, content in files.items():
                info=zipfile.ZipInfo(filename,(2026,10,30,0,0,0));info.compress_type=zipfile.ZIP_DEFLATED
                z.writestr(info,content.encode('utf-8-sig' if filename.endswith('.csv') else 'utf-8'))
        path=output/name;path.write_bytes(stream.getvalue())
        (output/(name+'.sha256')).write_text(digest(path.read_bytes())+'  '+name+'\n',encoding='ascii')
        store=SimpleNamespace(units=units,units_by_id=unit_by_id,
            report=SimpleNamespace(build_id=graph_fp,retrieval_unit_fingerprint=unit_fp,corpus_snapshot_as_of=DATE))
        imported, report=import_candidate(path,store,graph_fp)
        if report['duplicate_question_texts'] or report['query_after_snapshot_ids']:
            raise ValueError('Duplicate question or date beyond the snapshot')
        validation = {'questions':len(imported),'positive_retrieval_questions':report['retrieval_questions'],
            'evidence_records':len(catalog),'mapping_status':report['mapping_status'],
            'legal_approval':report['legal_approval'],'missing_gold_units':0,
            'model_inference':'NOT_RUN','all_annotated_evidence_present':True,
            'no_claim_of_complete_legal_domain':True,'artifact_build_id':graph_fp,
            'retrieval_unit_fingerprint':unit_fp,'gold_zip_sha256':digest(path.read_bytes()),
            'original_zip_preserved':digest(original_path.read_bytes())==manifest['source_parent_zip_sha256']}
        validations[name]=validation
        (output/(name.removesuffix('.zip')+'_validation.json')).write_text(dumps(report),encoding='utf-8')
    (output/'validation_summary.json').write_text(dumps(validations),encoding='utf-8')
    (output/'change_log.json').write_text(dumps(change_log),encoding='utf-8')
    write_workbook(output/'Gold_v0_2_review.xlsx',rows+extra,change_log,unit_by_id)
    print(dumps(validations))


def csv_text(rows):
    stream=io.StringIO(newline='')
    fields=['question_id','category','question','applicable_date','expected_status','gold_answer','gold_evidence_locators','missing_facts','annotation_note']
    writer=csv.DictWriter(stream,fieldnames=fields,extrasaction='ignore');writer.writeheader()
    for r in rows:
        writer.writerow({k:' | '.join(r[k]) if isinstance(r.get(k),list) else r.get(k,'') for k in fields})
    return stream.getvalue()


def write_workbook(path,rows,changes,units):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb=Workbook();ws=wb.active;ws.title='Cau_hoi_120'
    ws.append(['ID','Nhóm','Câu hỏi','Ngày áp dụng','Nhãn','Đáp án tham khảo','Căn cứ','Thiếu dữ kiện','Ghi chú','Người duyệt pháp lý','Kết quả duyệt'])
    for r in rows:
        ws.append([r['question_id'],r['category'],r['question'],r['applicable_date'],r['expected_status'],r['gold_answer'],'; '.join(r['gold_evidence_locators']),'; '.join(r['missing_facts']),r.get('annotation_note'),'','CHƯA DUYỆT'])
    evidence=wb.create_sheet('Can_cu');evidence.append(['Evidence ID','Điều','Khoản','Điểm','Nguyên văn trong artifact','Ngữ cảnh cha','URL nguồn'])
    for uid in sorted({i for r in rows for i in r['gold_evidence_ids']}):
        u=units[uid];evidence.append([uid,u['article_number'],u['clause_number'],u['point_number'],u['source_text'],u.get('ancestor_context'),u['source_url']])
    log=wb.create_sheet('Thay_doi');log.append(['ID','Trường sửa','Lý do','Trước','Sau'])
    for c in changes:log.append([c['question_id'],', '.join(c['changed_fields']),c['reason'],dumps(c['before']),dumps(c['after'])])
    note=wb.create_sheet('Huong_dan');note.append(['Lưu ý'])
    for text in ['Bản mặc định 100 câu: 88 câu có căn cứ, 12 câu cần hỏi thêm. 20 mã MR là phần mở rộng tùy chọn.',
        'Ngày đo 15/09/2026, không phải xác nhận luật hiện hành ở mọi ngày sau đó.',
        'Có căn cứ trong corpus không bảo đảm model luôn truy xuất/suy luận đúng; cần chạy benchmark inference riêng.',
        'Tất cả nhãn vẫn DRAFT, cần người có chuyên môn duyệt. Không coi kết quả máy kiểm là duyệt pháp lý.',
        'Không nhập đáp án/evidence ID vào câu hỏi hoặc dữ liệu retrieval. Không so điểm mới/cũ như cùng một tập Gold.',
        'Không xóa câu khó chỉ vì hệ thống trả lời sai. Ba câu tương lai giữ riêng để kiểm tra từ chối.']:
        note.append([text])
    for sheet in wb:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='215E83')
        for column in sheet.columns:
            letter=column[0].column_letter
            sheet.column_dimensions[letter].width=min(65,max(14,len(str(column[0].value))+5))
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                # Never execute spreadsheet input as formula when exported.
                if isinstance(cell.value,str) and cell.value[:1] in '=+@':cell.value="'"+cell.value
                cell.alignment=Alignment(vertical='top',wrap_text=True)
        for number in range(2,sheet.max_row+1):sheet.row_dimensions[number].height=75
    wb.save(path)


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--artifact',type=Path,required=True)
    parser.add_argument('--project',type=Path,default=Path(__file__).resolve().parents[2])
    parser.add_argument('--output',type=Path,required=True)
    args=parser.parse_args()
    build(args.project.resolve(),args.artifact.resolve(),args.output.resolve())
